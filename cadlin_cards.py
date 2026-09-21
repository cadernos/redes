#!/usr/bin/env python3
"""
cadlin_cards.py — Gera os cards de divulgação da Cadernos de Linguística (CadLin)
a partir dos templates oficiais do Estúdio Guayabo (Feed 4:5, R02), em TODAS as cores.

Uso:
    python3 cadlin_cards.py --spec spec.json --templates <pasta Feed_4x5_Meta> \
        --fonts <pasta _Fontes_Instalar_1o> --out <pasta de saída>

Saída: <out>/<slug>/  com
    PNG/<Cor>/<slug>_<card>_<modelo>_<PT|EN>_<Cor>_<fundo>.png   (1080x1350)
    PPTX/CadLin_<slug>_<Cor>.pptx                                  (editável)
    PAINEL_<Cor>.jpg                                               (miniaturas p/ escolha)
    RELATORIO.txt                                                  (checagens e pendências)
e um ZIP <out>/<slug>_cards.zip com tudo.

Marcação de texto no spec:
    {{termo}}  -> ênfase no estilo do próprio modelo (itálico nos títulos, negrito
                  na massa de texto, grifo colorido nas citações e respostas)
    **termo**  -> negrito forçado (usado nas editorias: "Leia em **Cadernos de Linguística**")
Colchetes comuns, como "[...]" em citações, são texto literal.
"""
import unicodedata, argparse, copy, glob, json, os, re, shutil, subprocess, sys, zipfile

from lxml import etree
from pptx import Presentation
from PIL import Image, ImageFont

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS = {"a": A}
EMU_PT = 12700

CORES = ["Azul", "Verde", "Amarelo", "Rosa", "Lilas"]
DECK = {c: f"CadLin_Feed_4x5_{c}_R02.pptx" for c in CORES}
DECK["Neutro"] = "CadLin_Feed_4x5_zNeutro_R02.pptx"

# Modelo -> slides-fonte (numeração da Tabela de Modelos; deck zNeutro tem numeração própria)
MODELOS = {
    "1A": {"cor": [1], "neutro": [1]},
    "1B": {"cor": [2], "neutro": [2]},
    "1C": {"cor": [3], "neutro": [3]},
    "1D": {"cor": [4], "neutro": [4]},
    "1E": {"cor": [5], "neutro": [5]},
    "1F": {"cor": [6], "neutro": [6]},
    "2A": {"cor": [7, 8]},
    "3":  {"cor": [9]},
    "4A": {"cor": [10, 11]},
    "6A": {"cor": [14, 15]},
    "7A": {"cor": [16], "neutro": [7]},
    "7B": {"cor": [17, 18]},
    "8A": {"cor": [19, 20]},
    "9A": {"cor": [21]},
    "9C": {"cor": [23]},
    "10A": {"cor": [24, 25], "neutro": [8]},
    "10B": {"cor": [26, 27]},
    "13A": {"cor": [32, 33], "neutro": [11]},
    "14A": {"cor": [34, 35], "neutro": [12]},
}

LIMITES = {  # caracteres (sem marcação) — do prompt editorial
    ("1C", "titulo"): 35, ("1A", "titulo"): 35, ("1E", "titulo"): 80,
    ("2A", "subtitulo"): 40, ("2A", "texto"): 270,
    ("3", "legenda"): 30, ("6A", "pergunta"): 67, ("6A", "resposta"): 320,
    ("13A", "legenda"): 80, ("14A", "subtitulo"): 30, ("14A", "texto"): 300,
    ("9A", "texto"): 300, ("10A", "minibio"): 350, ("10B", "minibio"): 350, ("8A", "citacao"): 300, ("7B", "citacao"): 100,
}

FONT_DIR = None


# ---------------------------------------------------------------- utilitários XML
def limpar(t):
    return re.sub(r"\*\*|\{\{|\}\}", "", t or "")


def rpr_sig(rpr):
    if rpr is None:
        return ("", "", False, "")
    lat = rpr.find("a:latin", NS)
    return (rpr.get("b", "0"), rpr.get("i", "0"),
            rpr.find("a:highlight", NS) is not None,
            lat.get("typeface") if lat is not None else "")


def estilos_do_paragrafo(p):
    """Retorna (rPr normal, rPr de ênfase) a partir dos runs do template."""
    runs = p.findall("a:r", NS)
    if not runs:
        return None, None
    normal = runs[0].find("a:rPr", NS)
    base = rpr_sig(normal)
    enf = None
    for r in runs:
        rp = r.find("a:rPr", NS)
        if rpr_sig(rp) != base:
            enf = rp
            break
    if enf is None:  # modelo sem ênfase própria: usa itálico
        enf = copy.deepcopy(normal)
        enf.set("i", "1")
    return normal, enf


def tokens(txt):
    """Divide texto em (trecho, estilo) com estilo em {n, e, b}."""
    out = []
    for part in re.split(r"(\{\{.*?\}\}|\*\*.*?\*\*)", txt):
        if not part:
            continue
        if part.startswith("{{"):
            out.append((part[2:-2], "e"))
        elif part.startswith("**"):
            out.append((part[2:-2], "b"))
        else:
            out.append((part, "n"))
    return out


def preencher_paragrafo(p, txt, normal=None, enf=None):
    if normal is None:
        normal, enf = estilos_do_paragrafo(p)
    if normal is None:
        normal = p.find("a:endParaRPr", NS)
        normal = copy.deepcopy(normal) if normal is not None else etree.SubElement(p, f"{{{A}}}rPr")
        normal.tag = f"{{{A}}}rPr"
        enf = copy.deepcopy(normal); enf.set("i", "1")
    for r in p.findall("a:r", NS) + p.findall("a:br", NS) + p.findall("a:fld", NS):
        p.remove(r)
    end = p.find("a:endParaRPr", NS)
    for trecho, est in tokens(txt):
        r = etree.Element(f"{{{A}}}r")
        rp = copy.deepcopy(enf if est == "e" else normal)
        rp.tag = f"{{{A}}}rPr"
        if est == "b":
            rp.set("b", "1")
        for k in ("err", "dirty"):
            rp.attrib.pop(k, None)
        r.append(rp)
        t = etree.SubElement(r, f"{{{A}}}t"); t.text = trecho
        if trecho != trecho.strip():
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        if end is not None:
            end.addprevious(r)
        else:
            p.append(r)


def preencher_txbody(txbody, mapa):
    """mapa: {indice_do_paragrafo_no_template: texto | [textos]}.
    Parágrafos com texto no template e ausentes do mapa são removidos; vazios ficam."""
    ps = txbody.findall("a:p", NS)
    for i, p in enumerate(ps):
        if i in mapa:
            val = mapa[i]
            vals = val if isinstance(val, list) else [val]
            normal, enf = estilos_do_paragrafo(p)
            preencher_paragrafo(p, vals[0], normal, enf)
            ant = p
            for extra in vals[1:]:
                novo = copy.deepcopy(p)
                preencher_paragrafo(novo, extra, normal, enf)
                ant.addnext(novo); ant = novo
        elif p.findall("a:r", NS):
            txbody.remove(p)


def texto_shape(sh):
    return sh.text_frame.text.strip() if sh.has_text_frame else ""


def nfc(t):
    return unicodedata.normalize("NFC", t).lower()


def achar(slide, prefixo):
    for sh in slide.shapes:
        if nfc(texto_shape(sh)).startswith(nfc(prefixo)):
            return sh
    return None


def achar_tabela(slide):
    for sh in slide.shapes:
        if sh.has_table:
            return sh
    return None


def txb(sh):
    return sh._element.find(".//p:txBody", {"p": "http://schemas.openxmlformats.org/presentationml/2006/main"})


# ---------------------------------------------------------------- duplicação de slides
def duplicar(prs, src):
    novo = prs.slides.add_slide(src.slide_layout)
    P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
    tree_n = novo.shapes._spTree
    for el in list(tree_n):
        tree_n.remove(el)
    for el in src.shapes._spTree:
        tree_n.append(copy.deepcopy(el))
    bg = src._element.cSld.find(P + "bg")
    if bg is not None:
        novo._element.cSld.insert(0, copy.deepcopy(bg))
    mapa = {}
    for rid, rel in src.part.rels.items():
        if rel.reltype.endswith("/slideLayout") or rel.reltype.endswith("/notesSlide"):
            continue
        if rel.is_external:
            nid = novo.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        else:
            nid = novo.part.rels.get_or_add(rel.reltype, rel.target_part)
        mapa[rid] = nid
    for el in novo._element.iter():
        for att in list(el.attrib):
            if att.startswith(f"{{{R}}}") and el.get(att) in mapa:
                el.set(att, mapa[el.get(att)])
    return novo


def remover_slide(prs, slide):
    lst = prs.slides._sldIdLst
    for sid in list(lst):
        if prs.part.related_part(sid.rId) is slide.part:
            prs.part.drop_rel(sid.rId); lst.remove(sid); return


# ---------------------------------------------------------------- etiquetas (7B)
def largura_texto(txt, pt=30, spc_pt=2.4, face="RedHatDisplay-Medium.ttf"):
    f = ImageFont.truetype(os.path.join(FONT_DIR, face), size=pt * 10)
    w = f.getlength(txt.upper()) / 10
    return int((w + spc_pt * len(txt)) * EMU_PT)


def fluxo_etiquetas(slide, palavras):
    prefixos = ["palavra-chave", "Lorem ipsum dolor", "amet", "ipsum dolor"]
    shapes = [achar(slide, p) for p in prefixos]
    shapes = [s for s in shapes if s is not None]
    x0, y0, gap, xmax = 1026000, 3657600, 242046, 1026000 + 11664000
    alt = shapes[0].height
    dy = 4746961 - 3657600
    x, y = x0, y0
    for sh, pal in zip(shapes, palavras):
        preencher_txbody(txb(sh), {0: pal})
        w = int(largura_texto(pal) * 1.12) + 432000 * 2 + 150000
        bp = txb(sh).find("a:bodyPr", NS)
        bp.set("wrap", "none")
        for fit in list(bp):
            bp.remove(fit)
        etree.SubElement(bp, f"{{{A}}}noAutofit")
        if x + w > xmax and x > x0:
            x, y = x0, y + dy
        sh.left, sh.top, sh.width, sh.height = x, y, w, alt
        x += w + gap
    for sh in shapes[len(palavras):]:
        sh._element.getparent().remove(sh._element)
    return y  # topo da última linha


# ---------------------------------------------------------------- preenchimento por modelo
def preencher(slide, modelo, c, relatorio, rotulo):
    def ed():
        sh = achar(slide, "Editoria") or achar(slide, "Revisão aberta") or achar(slide, "dado")
        if sh is not None and c.get("editoria"):
            preencher_txbody(txb(sh), {0: c["editoria"]})
        return sh

    m = modelo
    if m in ("1A", "1B", "1C", "1D", "1E", "1F"):
        ed()
        alvo = achar(slide, "Título") or achar(slide, "título")
        preencher_txbody(txb(alvo), {0: c["titulo"]})
        cta = achar(slide, "cta")
        if cta is not None:
            preencher_txbody(txb(cta), {0: c.get("cta", "Link na bio")})
    elif m == "2A":
        ed()
        sh = achar(slide, "Subtítulo") or achar(slide, "SUBTÍTULO")
        preencher_txbody(txb(sh), {0: c["subtitulo"], 2: c["texto"]})
    elif m == "3":
        ed()
        preencher_txbody(txb(achar(slide, "345")), {0: c["numero"]})
        preencher_txbody(txb(achar(slide, "legenda")), {0: c["legenda"]})
        sh = achar(slide, "Título de artigo")
        preencher_txbody(txb(sh), {1: c["titulo"], 2: c["referencia"]})
    elif m == "4A":
        ed()
        preencher_txbody(txb(achar(slide, "Título de artigo")), {0: c["titulo"]})
        tb = achar_tabela(slide).table
        autores = c["autores"]
        linhas = list(tb._tbl.tr_lst)
        for i, tr in enumerate(linhas):
            if i >= len(autores):
                tb._tbl.remove(tr)
        for i, (nome, afil) in enumerate(autores[:len(linhas)]):
            cel = tb.cell(i, 1)
            preencher_txbody(cel._tc.txBody, {0: nome, 1: afil})
        if len(autores) > len(linhas):
            relatorio.append(f"{rotulo}: template comporta {len(linhas)} autores; {len(autores)} informados.")
    elif m == "6A":
        ed()
        tb = achar_tabela(slide).table
        preencher_txbody(tb.cell(0, 0)._tc.txBody, {0: c["pergunta"]})
        preencher_txbody(tb.cell(1, 0)._tc.txBody, {0: c["resposta"]})
    elif m == "7A":
        ed()
        preencher_txbody(txb(achar(slide, "“Citação")), {0: c["citacao"], 1: c["referencia"]})
    elif m == "7B":
        ed()
        fluxo_etiquetas(slide, c["palavras"])
        preencher_txbody(txb(achar(slide, "“Citação")), {0: c["citacao"], 1: c["referencia"]})
    elif m == "8A":
        ed()
        tb = achar_tabela(slide).table
        preencher_txbody(tb.cell(0, 0)._tc.txBody, {0: c["citacao"]})
        preencher_txbody(tb.cell(1, 0)._tc.txBody, {0: c["referencia"]})
    elif m == "9A":
        ed()
        preencher_txbody(txb(achar(slide, "trecho de")), {0: c["cabecalho"], 1: c["autor"]})
        corpo = txb(achar(slide, "Parecer"))
        corpo.find("a:bodyPr", NS).set("anchor", "t")
        preencher_txbody(corpo, {0: c["texto"]})
    elif m == "9C":
        ed()
        mapa = {0: "1", 1: c["nota"]}
        mapa[3] = c.get("enviado", ""); mapa[4] = c.get("publicado", "")
        preencher_txbody(txb(achar(slide, "1")), mapa)
    elif m in ("10A", "10B"):
        ed()
        preencher_txbody(txb(achar(slide, "Nome da pessoa")), {0: c["nome"], 2: c["minibio"]})
        if m == "10A":
            img = c.get("imagem")
            ph = [s for s in slide.placeholders if s.placeholder_format.type is not None]
            if img and os.path.exists(img) and ph:
                ph[0].insert_picture(img)
            else:
                relatorio.append(f"{rotulo}: sem retrato — PNG não exportado; inserir foto no PPTX ou usar 10B.")
                return False
    elif m in ("13A", "14A"):
        if m == "14A":
            ed()
            sh = achar(slide, "subtítulo")
            preencher_txbody(txb(sh), {0: c["subtitulo"], 2: c["texto"]})
        else:
            preencher_txbody(txb(achar(slide, "Breve legenda")), {0: c["legenda"]})
        img = c.get("imagem")
        ph = [s for s in slide.placeholders if s.placeholder_format.type is not None]
        if img and os.path.exists(img) and ph:
            ph[0].insert_picture(img)
        else:
            relatorio.append(f"{rotulo}: sem imagem — PNG não exportado; inserir imagem no PPTX.")
            return False
    else:
        raise ValueError(f"Modelo não suportado: {m}")
    return True


# ---------------------------------------------------------------- render
def soffice_pdf(pptx, outdir):
    wrapper = "/mnt/skills/public/pptx/scripts/office/soffice.py"
    cmd = ([sys.executable, wrapper] if os.path.exists(wrapper) else ["soffice"]) + \
          ["--headless", "--convert-to", "pdf", "--outdir", outdir, pptx]
    subprocess.run(cmd, check=True, capture_output=True, timeout=600)
    return os.path.join(outdir, os.path.splitext(os.path.basename(pptx))[0] + ".pdf")


def fundo(png):
    im = Image.open(png).convert("RGB")
    r, g, b = im.getpixel((15, im.height // 2))
    return "fundo-neutro" if abs(r - 0xFA) < 10 and abs(g - 0xF6) < 10 and abs(b - 0xF4) < 10 else "fundo-colorido"


def painel(pngs, destino, titulo):
    if not pngs:
        return
    ims = [Image.open(p).convert("RGB").resize((270, 338)) for p in pngs]
    cols = 8; rows = (len(ims) + cols - 1) // cols
    g = Image.new("RGB", (cols * 280, rows * 348), (16, 20, 22))
    for i, im in enumerate(ims):
        g.paste(im, ((i % cols) * 280 + 5, (i // cols) * 348 + 5))
    g.save(destino, quality=85)


# ---------------------------------------------------------------- principal
def main():
    global FONT_DIR
    ap = argparse.ArgumentParser()
    ap.add_argument("--spec", required=True)
    ap.add_argument("--templates", required=True)
    ap.add_argument("--fonts", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--cores", default=",".join(CORES + ["Neutro"]))
    a = ap.parse_args()
    FONT_DIR = a.fonts

    fdir = os.path.expanduser("~/.fonts"); os.makedirs(fdir, exist_ok=True)
    for f in glob.glob(os.path.join(a.fonts, "*.ttf")):
        shutil.copy(f, fdir)
    subprocess.run(["fc-cache", "-f"], capture_output=True)

    spec = json.load(open(a.spec, encoding="utf-8"))
    slug = spec["slug"]
    base = os.path.join(a.out, slug)
    shutil.rmtree(base, ignore_errors=True)
    os.makedirs(os.path.join(base, "PPTX")); os.makedirs(os.path.join(base, "_tmp"))
    rel = []

    # checagem de limites
    for card in spec["cards"]:
        for lang in ("pt", "en"):
            for campo, v in card.get(lang, {}).items():
                lim = LIMITES.get((card["modelo"], campo))
                if lim and isinstance(v, str) and len(limpar(v)) > lim:
                    rel.append(f"{card['id']} {lang.upper()} {campo}: {len(limpar(v))} caracteres (limite {lim}).")

    for cor in a.cores.split(","):
        prs = Presentation(os.path.join(a.templates, DECK[cor]))
        originais = list(prs.slides)
        saida = []  # (slide, rotulo, exportar)
        for card in spec["cards"]:
            m = card["modelo"]
            fontes = MODELOS[m].get("neutro" if cor == "Neutro" else "cor", [])
            for n in fontes:
                for lang in ("pt", "en"):
                    if lang not in card:
                        continue
                    img = card.get("imagem") or card.get("pt", {}).get("imagem") or card.get("en", {}).get("imagem")
                    if img and "imagem" not in card[lang]:
                        card[lang]["imagem"] = img
                    s = duplicar(prs, originais[n - 1])
                    rot = f"{slug}_{card['id']}_M{m}-{n}_{lang.upper()}_{cor}"
                    ok = preencher(s, m, card[lang], rel, rot)
                    saida.append((s, rot, ok))
        if not saida:
            continue
        for s in originais:
            remover_slide(prs, s)
        pptx = os.path.join(base, "PPTX", f"CadLin_{slug}_{cor}.pptx")
        prs.save(pptx)

        pdf = soffice_pdf(pptx, os.path.join(base, "_tmp"))
        pdir = os.path.join(base, "PNG", cor); os.makedirs(pdir, exist_ok=True)
        pref = os.path.join(base, "_tmp", cor)
        subprocess.run(["pdftoppm", "-png", "-scale-to-x", "1080", "-scale-to-y", "1350", pdf, pref], check=True)
        pags = sorted(glob.glob(pref + "-*.png"), key=lambda x: int(re.findall(r"-(\d+)\.png$", x)[0]))
        feitos = []
        for (s, rot, ok), png in zip(saida, pags):
            if not ok:
                continue
            nome = f"{rot}_{fundo(png)}.png"
            shutil.move(png, os.path.join(pdir, nome)); feitos.append(os.path.join(pdir, nome))
        painel(feitos, os.path.join(base, f"PAINEL_{cor}.jpg"), cor)
        print(f"{cor}: {len(feitos)} PNG")

    shutil.rmtree(os.path.join(base, "_tmp"), ignore_errors=True)
    with open(os.path.join(base, "RELATORIO.txt"), "w", encoding="utf-8") as f:
        f.write("\n".join(rel) if rel else "Sem alertas.")
    zp = os.path.join(a.out, f"{slug}_cards.zip")
    with zipfile.ZipFile(zp, "w", zipfile.ZIP_DEFLATED) as z:
        for root, _, files in os.walk(base):
            for fn in files:
                full = os.path.join(root, fn)
                z.write(full, os.path.relpath(full, a.out))
    print("ZIP:", zp); print("\n".join(rel))


if __name__ == "__main__":
    main()
